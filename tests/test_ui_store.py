# -*- coding: utf-8 -*-
"""اختبارات حفظ أمر العمل وفتحه من الواجهة، وإدارة الأسعار (ق-٦١، ق-٦٢).

**الحارس الأهمّ:** أن كل حقل في اللوحة يُستعاد كما كان. فحقل يُضاف إلى `content()`
ويُنسى في `_load()` يجعل الفتح يعطي رقماً أقلّ **بصمت** — بلا خطأ ولا رسالة.
والاختبار هنا يملأ الكائن من `dataclasses.fields` لا من قائمة مكتوبة، فحقل جديد
يدخل الحراسة بلا تعديل هنا.
"""

import pytest

pytest.importorskip("PyQt6", reason="PyQt6 غير مثبَّت")

from datetime import date  # noqa: E402

from engine import load_catalog  # noqa: E402
from engine.types import (  # noqa: E402
    Equipment,
    Network11kV,
    Network33kV,
    NetworkLV,
    Project,
    Segment,
    SegmentKind,
    Underground11kV,
    Underground33kV,
)
from engine.workorder import WorkOrder  # noqa: E402
from ui.main_window import MainWindow  # noqa: E402

from .test_store import populate  # noqa: E402

UI_BOUNDS = {
    # حقول تحدّها الواجهة بمدى أو بدقّة — تُملأ بقيمة صالحة تبقى **غير افتراضية**
    "waste_pct": 0.15,          # الحقل نسبة مئوية 0..100، والمحرك يخزّنها كسراً
    "span_m": 30.0,
    "tension_span_m": 150.0,
    "drum_length_m": 400.0,
}


def ui_populate(cls):
    """كائن كل حقوله غير افتراضية **وكلها يمكن للواجهة تمثيلها**.

    الحقول العشرية التي حقلها في الواجهة عدد صحيح تُقرَّب إلى صحيح، وإلا فشلت
    المقارنة لسبب لا علاقة له بالحفظ. والحارس يبقى كاملاً: كل حقل يظلّ مختلفاً
    عن قيمته الافتراضية، ويتحقّق من ذلك اختبار مستقلّ.
    """
    obj = populate(cls, seed=3)
    for name, value in UI_BOUNDS.items():
        if hasattr(obj, name):
            setattr(obj, name, value)
    for name, value in vars(obj).items():
        if isinstance(value, float) and name not in UI_BOUNDS:
            setattr(obj, name, float(int(value)) or 3.0)
    return obj

KIND_OF = {
    Network11kV: SegmentKind.HV11,
    Network33kV: SegmentKind.HV33,
    NetworkLV: SegmentKind.LV,
    Equipment: SegmentKind.EQUIPMENT,
    Underground11kV: SegmentKind.UG11,
    Underground33kV: SegmentKind.UG33,
}


@pytest.fixture
def window(qapp):
    return MainWindow(load_catalog())


def full_project() -> Project:
    """مشروع فيه مقطع من كل نوع، كل حقوله غير افتراضية."""
    return Project(
        "مشروع كامل",
        [Segment(f"مقطع {i}", ui_populate(cls)) for i, cls in enumerate(KIND_OF)],
        street_crossing_secondary_m=12.0,
        street_crossing_secondary_feeders=3,
        street_crossing_main_m=8.0,
        street_crossing_main_feeders=2,
    )


# ═════════════ ١. حارس الحقول: اللوحة تُعيد ما حُمّل إليها ═════════════


@pytest.mark.parametrize("cls", list(KIND_OF), ids=lambda c: c.__name__)
def test_a_panel_returns_exactly_what_was_loaded_into_it(window, cls):
    """`load` ثم `content` = الأصل. أي حقل منسيّ في `_load` يفشل هنا."""
    original = ui_populate(cls)
    editor = window.add_segment(KIND_OF[cls])
    editor.load(original)
    assert editor.content() == original


def test_the_guard_would_catch_a_forgotten_field(window):
    """يثبت أن الحارس أعلاه ليس فارغاً: القيم المحمَّلة غير الافتراضية فعلاً."""
    original = ui_populate(Network11kV)
    editor = window.add_segment(SegmentKind.HV11)
    default = editor.content()
    differing = [f for f in vars(original) if getattr(original, f) != getattr(default, f)]
    assert len(differing) == len(vars(original)), differing


def test_loading_emits_one_change_not_one_per_field(window):
    """الفتح يُعيد الحساب مرّة لا مرّةً لكل حقل — وإلا تجمّد فتح مشروع كبير."""
    editor = window.add_segment(SegmentKind.HV11)
    calls = []
    editor.changed.connect(lambda: calls.append(1))
    editor.load(ui_populate(Network11kV))
    assert len(calls) == 1


def test_an_out_of_range_choice_is_refused_not_silently_ignored(window):
    """قيمة لا توجد في القائمة تُعطّل صراحةً — لا تترك القائمة على خيارها الأول."""
    from ui.panels import _select

    editor = window.add_segment(SegmentKind.HV11)
    with pytest.raises(ValueError, match="لا توجد في القائمة"):
        _select(editor.circuit, "ثلاثية")


# ═════════════ ٢. الحفظ والفتح من النافذة ═════════════


def test_the_window_round_trips_a_whole_work_order(window, tmp_path):
    """أمر عمل كامل: يُحفظ، ثم يُفتح في نافذة أخرى، فيتطابقان."""
    project = full_project()
    window.segments.load(project)
    order = WorkOrder(number="45", order_date=date(2026, 8, 31),
                      classification="توسعات", project_name="مشروع كامل",
                      duration="90 يوم", start_date=None, notes="ملاحظة")
    order.staff[0].count, order.staff[0].days = 2, 90
    window.order_panel.load(order)

    path = window.save_to(tmp_path / "أمر")
    before = window.result["الكلفة_الكلية"]

    other = MainWindow(load_catalog())
    other.load_from(path)
    assert other.project() == window.project()
    assert other.order_panel.order() == order
    assert other.result["الكلفة_الكلية"] == before


def test_saving_records_the_price_version_and_opening_restores_it(window, tmp_path):
    """أمر عمل أُنشئ بأسعار نسخةٍ يُعاد فتحه بها لا بأحدث نسخة (ق-٤٠)."""
    window.segments.load(Project("م", [Segment("أ", Network11kV(poles_lattice=3))]))
    path = window.save_to(tmp_path / "أمر")

    _, _, version = __import__("engine.store", fromlist=["load"]).load(path)
    assert version == window.version

    other = MainWindow(load_catalog())
    other.load_from(path)
    assert other.version == version


def test_the_title_names_the_file_and_the_price_version(window, tmp_path):
    """العنوان يفرّق أمر عمل عن آخر — ويُظهر أي أسعار تُحسب به."""
    assert "لم يُحفظ" in window.windowTitle()
    window.save_to(tmp_path / "أمر ٤٥")
    assert "أمر ٤٥.wo" in window.windowTitle()
    assert window.version in window.windowTitle()


def test_opening_replaces_the_previous_content_and_leaves_nothing_behind(
    window, tmp_path
):
    """فتح ملف ثانٍ لا يترك مقاطع الأول — وإلا تضاعفت الكميات بصمت."""
    window.segments.load(Project("أول", [Segment("أ", Network11kV(poles_lattice=9))]))
    first = window.save_to(tmp_path / "أول")

    window.segments.load(Project("ثانٍ", [
        Segment("ب", Network11kV(poles_lattice=1)),
        Segment("ج", Network33kV(poles_suspension=4)),
    ]))
    window.save_to(tmp_path / "ثانٍ")

    window.load_from(first)
    segments = window.project().segments
    assert len(segments) == 1
    assert segments[0].content.poles_lattice == 9


def test_a_new_order_empties_everything(window, tmp_path, monkeypatch):
    """«أمر عمل جديد» يبدأ من الصفر — بعد تأكيد، فالمُدخَل يضيع بلا رجعة."""
    monkeypatch.setattr(MainWindow, "_confirm", staticmethod(lambda *a: True))
    window.segments.load(full_project())
    window.save_to(tmp_path / "أمر")
    assert window.path is not None

    window.new_order()
    assert window.project().segments == []
    assert window.path is None
    assert window.order_panel.number.text() == ""


def test_declining_the_confirmation_keeps_the_work(window, monkeypatch):
    """الرفض يُبقي كل شيء — لا حذف جزئي."""
    monkeypatch.setattr(MainWindow, "_confirm", staticmethod(lambda *a: False))
    window.segments.load(full_project())
    count = len(window.project().segments)
    window.new_order()
    assert len(window.project().segments) == count


# ═════════════ ٣. الأسعار: التثبيت والتحديث بأمر صريح ═════════════


def test_switching_the_version_recalculates_the_cost(window, tmp_path, monkeypatch):
    """تغيير النسخة يغيّر الكلفة — وهو كلّ معنى تثبيت الأسعار."""
    import engine
    import engine.paths as paths
    from engine.prices import MATERIALS, PRICE, apply_edits, save_as_new_version
    import engine.prices as prices

    base = load_catalog()          # قبل التحويل: المجلد المؤقّت لا أسعار فيه
    monkeypatch.setattr(paths, "user_data_dir", lambda: tmp_path)
    monkeypatch.setattr(paths, "bundled_data_dir", lambda: tmp_path)
    monkeypatch.setattr(prices, "catalog_path", lambda v: tmp_path / f"catalog_{v}.json")
    save_as_new_version(base, "2000-01")
    dearer = apply_edits(base, [{
        "الباب": MATERIALS, "الاسم": "براكيت جنل 1.4 م مع الملحقات",
        "المفتاح": PRICE, "السعر": 1_000_000,
    }])
    save_as_new_version(dearer, "2000-02")

    win = MainWindow(engine.load_catalog("2000-01"), "2000-01")
    win.segments.load(Project("م", [Segment("أ", Network11kV(
        poles_lattice=9, extra_bracket_14=4))]))
    cheap = win.result["الكلفة_الكلية"]

    win.switch_version("2000-02")
    assert win.version == "2000-02"
    assert win.result["الكلفة_الكلية"] > cheap


def test_an_old_order_keeps_its_prices_when_a_newer_version_exists(
    window, tmp_path, monkeypatch
):
    """الحارس الجوهري: نسخة أحدث على القرص **لا تغيّر** كلفة أمر عمل قديم."""
    import engine
    import engine.paths as paths
    from engine.prices import MATERIALS, PRICE, apply_edits, save_as_new_version
    import engine.prices as prices

    base = load_catalog()          # قبل التحويل: المجلد المؤقّت لا أسعار فيه
    monkeypatch.setattr(paths, "user_data_dir", lambda: tmp_path)
    monkeypatch.setattr(paths, "bundled_data_dir", lambda: tmp_path)
    monkeypatch.setattr(prices, "catalog_path", lambda v: tmp_path / f"catalog_{v}.json")
    save_as_new_version(base, "2000-01")

    win = MainWindow(engine.load_catalog("2000-01"), "2000-01")
    win.segments.load(Project("م", [Segment("أ", Network11kV(
        poles_lattice=9, extra_bracket_14=4))]))
    original_cost = win.result["الكلفة_الكلية"]
    path = win.save_to(tmp_path / "قديم")

    # تُصدَر نسخة أغلى **بعد** حفظ أمر العمل
    save_as_new_version(apply_edits(base, [{
        "الباب": MATERIALS, "الاسم": "براكيت جنل 1.4 م مع الملحقات",
        "المفتاح": PRICE, "السعر": 1_000_000,
    }]), "2000-02")

    reopened = MainWindow(engine.load_catalog("2000-02"), "2000-02")
    reopened.load_from(path)
    assert reopened.version == "2000-01"
    assert reopened.result["الكلفة_الكلية"] == original_cost


# ═════════════ ٤. العمل على أكثر من حاسبة: نسخة أسعار غائبة ═════════════
#
# أمر عمل يحمل **اسم** نسخة الأسعار لا الأسعار نفسها. فمن ينشئ نسخة أسعار على
# حاسبة ويفتح أمر العمل على غيرها يجد النسخة غائبة. والرسالة يجب أن تقول **ما
# العمل**، لا أن تكتفي بـ«تعذّر الفتح».


def test_a_missing_price_version_says_what_to_do_about_it(window, tmp_path, monkeypatch):
    """الحارس: الرسالة الخاصّة تصل فعلاً.

    وكانت **لا تصل**: `FileNotFoundError` فرعٌ من `OSError`، وفرع `OSError` كان
    مكتوباً قبله، فيبتلعه ويبقى فرع الأسعار المفقودة ميتاً لا يُبلَّغ منه شيء.
    """
    import json

    from PyQt6.QtWidgets import QFileDialog, QMessageBox

    path = window.save_to(tmp_path / "أمر")
    data = json.loads(path.read_text(encoding="utf-8"))
    data["نسخة_الأسعار"] = "1999-01"          # نسخة لا وجود لها على هذه الحاسبة
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    shown = []
    monkeypatch.setattr(QFileDialog, "getOpenFileName",
                        staticmethod(lambda *a, **k: (str(path), "")))
    monkeypatch.setattr(QMessageBox, "critical",
                        staticmethod(lambda parent, title, text: shown.append((title, text))))

    window.open_order()

    assert shown, "لم تظهر أي رسالة"
    title, text = shown[0]
    assert title == "نسخة الأسعار مفقودة", title
    assert "1999-01" in text                  # تُسمّى النسخة الغائبة بعينها
    assert "catalog_" in text and "data" in text     # ويُذكر ما يُنسخ وإلى أين


def test_the_work_survives_a_failed_open(window, tmp_path, monkeypatch):
    """فتحٌ فاشل لا يمسح ما على الشاشة — وإلا ضاع عمل ساعة برسالة خطأ."""
    import json

    from PyQt6.QtWidgets import QFileDialog, QMessageBox

    window.segments.load(full_project())
    before = window.project()

    path = window.save_to(tmp_path / "أمر")
    data = json.loads(path.read_text(encoding="utf-8"))
    data["نسخة_الأسعار"] = "1999-01"
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    monkeypatch.setattr(QFileDialog, "getOpenFileName",
                        staticmethod(lambda *a, **k: (str(path), "")))
    monkeypatch.setattr(QMessageBox, "critical", staticmethod(lambda *a: None))

    window.open_order()
    assert window.project() == before


def test_the_prices_window_warns_that_an_edit_is_local_to_this_computer(qapp):
    """من يعمل على حاسبتين يجب أن يعرف — **قبل الاعتماد** — أنه يفرّق بينهما.

    فالنسخة الجديدة تُكتب على هذه الحاسبة وحدها، والبرنامج يبدأ بأحدث نسخة
    متاحة، فتصير أوامر العمل الجديدة هنا بأسعار غير أسعار الحاسبة الأخرى بلا
    ما ينبّه — إلا اسم النسخة في عنوان النافذة.
    """
    from PyQt6.QtWidgets import QLabel

    from ui.prices_window import PricesWindow

    window = PricesWindow(load_catalog(), "2026-08")
    hints = " ".join(label.text() for label in window.findChildren(QLabel))
    assert "أكثر من حاسبة" in hints
    assert "هذه الحاسبة وحدها" in hints


# ═════════════ ٥. تأشير المواد وخيارات قالب الأمانة تُحفظ وتُستعاد (ق-٧٦) ═════════════


def test_the_unavailable_ticks_survive_a_save_and_reopen(window, tmp_path):
    """التأشير جزء من أمر العمل — يُحفظ معه ولا يُعاد إدخاله عند كل طباعة."""
    window.segments.load(Project("م", [Segment("أ", Network11kV(poles_lattice=9))]))
    name = window._rows[0]["المادة"]
    window.unavailable = {name}
    path = window.save_to(tmp_path / "أمر")

    other = MainWindow(load_catalog())
    other.load_from(path)
    assert other.unavailable == {name}
    assert other.order().unavailable_materials == [name]


def test_the_amana_options_survive_a_save_and_reopen(window, tmp_path):
    window.order_panel.show_material_prices.setChecked(True)
    window.order_panel.preparation_committee.setValue(4)
    window.order_panel.audit_committee.setValue(3)
    path = window.save_to(tmp_path / "أمر")

    other = MainWindow(load_catalog())
    other.load_from(path)
    reopened = other.order()
    assert reopened.show_material_prices is True
    assert (reopened.preparation_committee, reopened.audit_committee) == (4, 3)


def test_a_new_order_clears_the_ticks(window, tmp_path, monkeypatch):
    """وإلا انتقل تأشير أمر عمل قديم إلى أمر عمل جديد بلا أن ينتبه أحد."""
    monkeypatch.setattr(MainWindow, "_confirm", staticmethod(lambda *a: True))
    window.segments.load(Project("م", [Segment("أ", Network11kV(poles_lattice=9))]))
    window.unavailable = {window._rows[0]["المادة"]}

    window.new_order()
    assert window.unavailable == set()


def test_the_excel_export_follows_the_selected_template(window, tmp_path):
    """زرّ الإكسل يُخرج ورقة القالب المعروض — لا ورقةً واحدة لكل القوالب (ق-٧٦)."""
    import openpyxl

    from printing.amana_sheet import SHEET
    from printing.spreadsheet import ORDER_SHEET

    window.segments.load(Project("م", [Segment("أ", Network11kV(poles_lattice=9))]))

    iso = window.write_order_xlsx(str(tmp_path / "إيزو"), template_key="iso")
    amana = window.write_order_xlsx(str(tmp_path / "أمانة"), template_key="amana")

    assert ORDER_SHEET in openpyxl.load_workbook(iso).sheetnames
    assert openpyxl.load_workbook(amana).sheetnames == [SHEET]
