# -*- coding: utf-8 -*-
"""ورقة إكسل لقالب «تنفيذ أمانة» — بمعادلات حيّة كنظيرها في ق-٧١ (ق-٧٦).

ورقة واحدة تحاكي المطبوع: المواد، ثم الأعمال المدنية، ثم الأعمال الكهربائية،
ثم المجاميع والملاحظات واللجنتان. وكل كلفة **معادلة** لا رقماً، فتعديل كمية أو
سعر يُحدّث سطره ومجاميعه داخل الإكسل.

## أسعار المواد تُكتب ويُخفى عرضها — ولا تُفرَّغ الخلايا

طلب المستخدم أن تكون خانتا «سعر المفرد» و«السعر الكلي» في جدول المواد **فارغتين
افتراضياً**، والورقة مع ذلك تحمل **مجموع كلفة المواد**.

ولو فُرِّغت الخانتان فعلاً لصار المجموع صفراً، أو — وهو أسوأ — رقماً جامداً لا
يساوي سطوره: ورقةٌ تُظهر مجموعاً لا تُنتجه أعمدتها.

فالحلّ أن **القيمة تبقى والعرض يُخفى**: خلايا أسعار المواد وحدها تأخذ التنسيق
`;;;` الذي يجعل إكسل لا يعرض محتواها على الشاشة ولا في الطباعة، وهي مع ذلك
أرقام ومعادلات حيّة تُغذّي المجاميع. وجدولا الأعمال المدنية والكهربائية في
الورقة نفسها **تظهر أسعارهما كاملة** — ولذلك كان الإخفاء **بتنسيق الخلية لا
بإخفاء العمود**: العمود واحد للجداول الثلاثة، فإخفاؤه كان سيمحو أسعار العمل
معه.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from engine.availability import unavailable_summary
from engine.workorder import WorkOrder

from .amana_form import (
    AUDIT_COMMITTEE,
    NOTES,
    PREPARATION_COMMITTEE,
    SIGNATURES_PER_ROW,
    SIGNATURE_LINE,
    header_line,
    printed_labour_name,
    printed_materials,
    printed_unit,
    split_labour,
    title_line,
)
from .spreadsheet import (
    _BOX,
    _CENTRE,
    _MONEY,
    _RIGHT,
    _TOTAL_FILL,
    _boxed,
    _table_head,
    _widths,
)

SHEET = "تنفيذ أمانة"

COLUMNS = ["ت", "التفاصيل", "الوحدة", "الكمية", "سعر المفرد د.ع", "السعر الكلي د.ع"]
LAST = len(COLUMNS)
"""عدد الأعمدة — وآخر عمود هو «السعر الكلي» (العمود F)."""

HIDDEN = ";;;"
"""تنسيق إكسل يُخفي عرض محتوى الخلية ويُبقي قيمتها — لأسعار المواد وحدها."""


def _merged(sheet, row: int, text: str, *, size: int = 11, bold: bool = False,
            align: str = "center") -> None:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=LAST)
    cell = sheet.cell(row, 1, text)
    cell.font = Font(bold=bold, size=size)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)


def _summary(sheet, row: int, label: str, value, *, bold: bool = True) -> None:
    """صفّ مجموع: عنوانه ممتدّ على خمسة أعمدة وقيمته في السادس."""
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=LAST - 1)
    cell = sheet.cell(row, 1, label)
    cell.font = Font(bold=bold)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    total = sheet.cell(row, LAST, value)
    total.font = Font(bold=True)
    total.number_format = _MONEY
    total.alignment = _CENTRE
    for column in range(1, LAST + 1):
        sheet.cell(row, column).fill = _TOTAL_FILL
        sheet.cell(row, column).border = _BOX


def _materials(sheet, result: dict, order: WorkOrder, row: int) -> dict:
    """جدول المواد. يعيد مواضع الصفوف التي تشير إليها المجاميع."""
    money = _MONEY if order.show_material_prices else HIDDEN
    _table_head(sheet, row, 1, COLUMNS)
    first = row + 1
    at = first
    materials = printed_materials(result)
    for index, material in enumerate(materials, start=1):
        sheet.cell(at, 1, index).alignment = _CENTRE
        sheet.cell(at, 2, material["المادة"]).alignment = _RIGHT
        sheet.cell(at, 3, material["الوحدة"]).alignment = _CENTRE
        sheet.cell(at, 4, material["الكمية"]).alignment = _CENTRE
        price = 0 if material["كمية_فقط"] else (material["سعر الوحدة"] or 0)
        sheet.cell(at, 5, price).number_format = money
        sheet.cell(at, 6, f"=D{at}*E{at}").number_format = money
        _boxed(sheet, at, 1, LAST)
        at += 1

    summary = unavailable_summary(materials, order.unavailable_materials)
    if summary["فقرات"]:
        numbers = "، ".join(str(n) for n in summary["فقرات"])
        _merged(sheet, at, f"المواد غير المتوفرة تشمل الفقرات ({numbers})")
        _boxed(sheet, at, 1, LAST)
        at += 1

    # **مجموع المؤشَّرة معادلةً** تجمع خلايا كلفتها بأعيانها، فتعديل كمية مادة
    # غير متوفرة يُحدّث كلفتها هنا كما يُحدّث المجموع الكلي.
    cells = [f"F{first + n - 1}" for n in summary["فقرات"]]
    unavailable_row = at
    _summary(sheet, at, "الكلفة التخمينية للمواد غير المتوفرة",
             f"={'+'.join(cells)}" if cells else 0, bold=False)
    at += 1

    total_row = at
    _summary(sheet, at, "الكلفة التخمينية الكلية للمواد",
             f"=SUM(F{first}:F{first + len(materials) - 1})" if materials else 0)
    return {"غير_المتوفرة": unavailable_row, "المجموع": total_row, "التالي": at + 2}


def _labour(sheet, title: str, lines: list, row: int, total_label: str) -> dict:
    """جدول أجور (مدني أو كهربائي). يعيد صفّ مجموعه والصفّ التالي."""
    # رأس العمود الثاني هو عنوان الجدول — كما في المطبوع تماماً
    _table_head(sheet, row, 1, [COLUMNS[0], title] + COLUMNS[2:])

    first = row + 1
    at = first
    for index, line in enumerate(lines, start=1):
        sheet.cell(at, 1, index).alignment = _CENTRE
        sheet.cell(at, 2, printed_labour_name(line.name)).alignment = _RIGHT
        sheet.cell(at, 3, printed_unit(line.unit)).alignment = _CENTRE
        sheet.cell(at, 4, line.qty).alignment = _CENTRE
        sheet.cell(at, 5, line.rate or 0).number_format = _MONEY
        sheet.cell(at, 6, f"=D{at}*E{at}").number_format = _MONEY
        _boxed(sheet, at, 1, LAST)
        at += 1
    if not lines:
        _boxed(sheet, at, 1, LAST)
        at += 1

    _summary(sheet, at, total_label,
             f"=SUM(F{first}:F{at - 1})" if lines else 0)
    return {"المجموع": at, "التالي": at + 2}


def _committee(sheet, row: int, title: str, slots: int) -> int:
    """عنوان اللجنة وتحته خطوط توقيع فارغة — بلا أسماء (بنصّ المستخدم)."""
    cell = sheet.cell(row, 1, title)
    cell.font = Font(bold=True, size=12)
    cell.alignment = Alignment(horizontal="right")
    at = row + 1
    for start in range(0, max(0, slots), SIGNATURES_PER_ROW):
        count = min(SIGNATURES_PER_ROW, slots - start)
        for offset in range(count):
            column = 1 + offset * 2
            sheet.merge_cells(start_row=at, start_column=column,
                              end_row=at, end_column=column + 1)
            sheet.cell(at, column, SIGNATURE_LINE).alignment = _CENTRE
        at += 1
    return at + 1


def build_workbook(order: WorkOrder, result: dict) -> Workbook:
    """المصنَّف كاملاً — ورقة واحدة تحاكي مطبوع «تنفيذ أمانة»."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet(SHEET)
    sheet.sheet_view.rightToLeft = True
    sheet.sheet_view.showGridLines = False
    _widths(sheet, {1: 6, 2: 46, 3: 12, 4: 12, 5: 16, 6: 18})

    _merged(sheet, 1, header_line(order), size=14, bold=True)
    _merged(sheet, 2, _plain_title(order), size=13, bold=True)

    places = _materials(sheet, result, order, row=4)
    civil, electrical = split_labour(result)
    civil_at = _labour(sheet, "الأعمال المدنية", civil, places["التالي"],
                       "مجموع اجور عمل الاعمال المدنية")
    electrical_at = _labour(sheet, "الأعمال الكهربائية", electrical,
                            civil_at["التالي"], "مجموع اجور عمل الاعمال الكهربائية")

    labour = f"(F{civil_at['المجموع']}+F{electrical_at['المجموع']})"
    row = electrical_at["التالي"]
    _summary(sheet, row, "مجموع اجور العمل (الاعمال الكهربائية + الاعمال المدنية)",
             f"={labour}", bold=False)
    _summary(sheet, row + 1, "مجموع أجور العمل + كلفة المواد غير المتوفرة",
             f"={labour}+F{places['غير_المتوفرة']}", bold=False)
    _summary(sheet, row + 2, "مجموع أجور العمل + كلفة المواد الكلية",
             f"={labour}+F{places['المجموع']}")

    row += 4
    for line in list(NOTES) + [
        note for note in (order.notes or "").splitlines() if note.strip()
    ]:
        _merged(sheet, row, f"• {line}", align="right")
        row += 1

    row = _committee(sheet, row + 1, PREPARATION_COMMITTEE,
                     order.preparation_committee)
    _committee(sheet, row, AUDIT_COMMITTEE, order.audit_committee)

    return workbook


def _plain_title(order: WorkOrder) -> str:
    """عنوان الورقة بلا تهريب HTML — الإكسل يكتب النصّ كما هو."""
    import html

    return html.unescape(title_line(order))


def write_xlsx(order: WorkOrder, result: dict, path: str) -> str:
    """يكتب المصنَّف إلى المسار ويعيده."""
    if not path.lower().endswith(".xlsx"):
        path += ".xlsx"
    build_workbook(order, result).save(path)
    return path
